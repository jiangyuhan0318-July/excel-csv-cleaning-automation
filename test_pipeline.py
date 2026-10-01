from datetime import datetime
from decimal import Decimal
import hashlib
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
import pandas as pd
from make_sample import generate
from pipeline import COLUMNS, DEFAULT_RULES, clean, normalize_row, read_input, run


def hashes(root):
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in root.rglob('*') if p.is_file()}


class CleaningTests(unittest.TestCase):
    def test_full_csv_and_excel(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            generate(root/'raw')
            before = hashes(root/'raw')
            csv_report = run(root/'raw/messy_orders.csv', root/'csv')
            excel_report = run(root/'raw/messy_orders.xlsx', root/'excel', sheet='Orders')
            self.assertEqual(csv_report, excel_report)
            self.assertEqual((csv_report['input_rows'], csv_report['cleaned_rows'], csv_report['duplicates_removed'], csv_report['rejected_rows']), (20,10,2,8))
            self.assertEqual(csv_report['normalized_cells'], 45)
            self.assertEqual(csv_report['amount_total'], '3150.99')
            self.assertEqual(csv_report['amount_mean'], '315.10')
            self.assertEqual(hashes(root/'raw'), before)
            self.assertEqual((root/'csv/cleaned.csv').read_bytes(), (root/'excel/cleaned.csv').read_bytes())
            frame = read_input(root/'csv/cleaned.csv')
            self.assertEqual(len(frame), 10)
            self.assertTrue(frame['order_id'].is_unique)
            self.assertEqual(sum(map(Decimal, frame['amount'])), Decimal('3150.99'))
            self.assertEqual(clean(frame)[2]['rejected_rows'], 0)
            self.assertEqual(clean(frame)[2]['normalized_cells'], 0)
            self.assertTrue(all(not normalize_row(row,DEFAULT_RULES)[1] for row in frame.to_dict('records')))
            audit = pd.read_csv(root/'csv/before_after.csv',keep_default_na=False)
            self.assertEqual(audit['decision'].value_counts().to_dict(), {'accepted':10,'rejected':8,'duplicate':2})
            self.assertEqual(audit['source_row'].tolist(), list(range(2,22)))
            wb = load_workbook(root/'csv/cleaned.xlsx',data_only=False)
            self.assertEqual(wb.sheetnames, ['Summary','Cleaned'])
            self.assertEqual((wb['Cleaned'].max_row, wb['Cleaned'].max_column),(11,7))
            self.assertIsInstance(wb['Cleaned']['E2'].value, datetime)
            self.assertIsInstance(wb['Cleaned']['F2'].value, (int,float))
            self.assertEqual(wb['Cleaned'].freeze_panes,'A2')
            self.assertEqual(wb['Cleaned'].auto_filter.ref,'A1:G11')
            self.assertEqual(wb['Summary']['B9'].value,3150.99)
            self.assertEqual(clean(read_input(root/'csv/cleaned.xlsx','Cleaned'))[2]['cleaned_rows'],10)
            self.assertEqual(run(root/'raw/messy_orders.csv',root/'again'),csv_report)
            with self.assertRaises(FileExistsError):
                run(root/'raw/messy_orders.csv',root/'csv')
            with self.assertRaises(ValueError):
                run(root/'raw/messy_orders.csv',root/'raw')

    def test_rules_and_boundaries(self):
        row = dict(zip(COLUMNS,['o1',' Test   Name ','','office','09/02/2026','0','paid']))
        normalized, errors = normalize_row(row,DEFAULT_RULES)
        self.assertEqual(errors,[])
        self.assertEqual(normalized['order_date'],'2026-09-02')
        self.assertEqual(normalized['amount'],'0.00')
        for amount in ('NaN','inf','1,20','10.001'):
            self.assertIn('invalid_amount',normalize_row({**row,'amount':amount},DEFAULT_RULES)[1])
        self.assertIn('unknown_status',normalize_row({**row,'status':'mystery'},DEFAULT_RULES)[1])
        self.assertIn('invalid_date',normalize_row({**row,'order_date':'2026-02-30'},DEFAULT_RULES)[1])
        rules = {**DEFAULT_RULES,'date_formats':['%d/%m/%Y']}
        self.assertEqual(normalize_row(row,rules)[0]['order_date'],'2026-02-09')
        cleaned, audit, report = clean(pd.DataFrame(columns=COLUMNS))
        self.assertTrue(cleaned.empty)
        self.assertEqual(report['amount_total'],'0.00')
        self.assertIsNone(report['amount_mean'])
        frame = pd.DataFrame([row,{**row,'amount':'20'},{**row,'amount':'20'}])
        self.assertEqual(clean(frame)[2]['cleaned_rows'],0)
        self.assertEqual(clean(frame)[2]['rejected_rows'],2)
        self.assertEqual(clean(frame)[2]['duplicates_removed'],1)

    def test_missing_schema(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp)/'bad.csv'
            source.write_text('order_id,amount\n001,20\n')
            with self.assertRaises(ValueError):
                run(source,Path(tmp)/'output')
            self.assertFalse((Path(tmp)/'output').exists())


if __name__ == '__main__':
    unittest.main()
