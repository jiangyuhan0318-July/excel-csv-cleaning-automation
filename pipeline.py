"""Reusable CSV/XLSX order cleaning with explicit rules and a row-level audit."""
import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = ['order_id', 'customer', 'email', 'category', 'order_date', 'amount', 'status']
DEFAULT_RULES = {
    'category_aliases': {'office': 'Office', 'office supplies': 'Office', 'tech': 'Technology', 'technology': 'Technology', 'furniture': 'Furniture'},
    'status_aliases': {'paid': 'Paid', 'complete': 'Paid', 'completed': 'Paid', 'pending': 'Pending', 'cancelled': 'Cancelled', 'canceled': 'Cancelled'},
    'date_formats': ['%Y-%m-%d', '%Y/%m/%d', '%m/%d/%Y', '%d-%b-%Y'],
}


def text(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    return str(value) if value is not None and not pd.isna(value) else ''


def read_input(path, sheet=0):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding='utf-8-sig')
    elif path.suffix.lower() == '.xlsx':
        frame = pd.read_excel(path, sheet_name=sheet, dtype=object, keep_default_na=False, engine='openpyxl')
    else:
        raise ValueError('Input must be .csv or .xlsx')
    frame.columns = [re.sub(r'\s+', '_', str(c).strip().lower()) for c in frame.columns]
    if list(frame.columns).count('order_id') != 1 or set(frame.columns) != set(COLUMNS) or len(frame.columns) != len(COLUMNS):
        raise ValueError('Expected exactly these columns: '+', '.join(COLUMNS))
    return frame[COLUMNS].map(text)


def normalize_row(raw, rules):
    row = {k: ' '.join(raw[k].split()) for k in COLUMNS}
    row['order_id'] = row['order_id'].upper()
    row['email'] = row['email'].lower()
    errors = []
    for col in ('order_id', 'customer', 'category', 'order_date', 'amount', 'status'):
        if not row[col]:
            errors.append('missing_'+col)
    if row['email'] and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', row['email']):
        errors.append('invalid_email')
    for col in ('category', 'status'):
        if row[col]:
            canonical = rules[col+'_aliases'].get(row[col].casefold())
            if canonical is None:
                errors.append('unknown_'+col)
            else:
                row[col] = canonical
    if row['order_date']:
        for fmt in rules['date_formats']:
            try:
                row['order_date'] = datetime.strptime(row['order_date'], fmt).date().isoformat()
                break
            except ValueError:
                pass
        else:
            errors.append('invalid_date')
    if row['amount']:
        amount = row['amount'].removeprefix('$')
        # US-style optional thousands separators, at most two decimal places.
        if not re.fullmatch(r'-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d{1,2})?', amount):
            errors.append('invalid_amount')
        else:
            try:
                number = Decimal(amount.replace(',', ''))
                if number < 0:
                    errors.append('negative_amount')
                else:
                    row['amount'] = f'{number:.2f}'
            except InvalidOperation:
                errors.append('invalid_amount')
    return row, errors


def clean(frame, rules=None):
    rules = rules or DEFAULT_RULES
    audit, seen = [], {}
    for source_row, raw in enumerate(frame.to_dict('records'), 2):
        row, errors = normalize_row(raw, rules)
        record = {'source_row': source_row, 'raw': raw, 'clean': row, 'decision': 'rejected' if errors else 'accepted', 'issues': errors}
        if not errors:
            key = tuple(row[c] for c in COLUMNS)
            if key in seen:
                record.update(decision='duplicate', issues=['duplicate_of_row_'+str(seen[key])])
            else:
                seen[key] = source_row
        audit.append(record)
    ids = Counter(r['clean']['order_id'] for r in audit if r['decision'] == 'accepted')
    conflicts = {key for key, count in ids.items() if count > 1}
    for record in audit:
        if record['decision'] == 'accepted' and record['clean']['order_id'] in conflicts:
            record.update(decision='rejected', issues=['conflicting_order_id'])
    accepted = [r['clean'] for r in audit if r['decision'] == 'accepted']
    total = sum((Decimal(r['amount']) for r in accepted), Decimal('0'))
    counts = Counter(r['decision'] for r in audit)
    report = {'input_rows': len(frame), 'cleaned_rows': len(accepted), 'duplicates_removed': counts['duplicate'],
              'rejected_rows': counts['rejected'], 'normalized_cells': sum(r['raw'][c] != r['clean'][c] for r in audit for c in COLUMNS),
              'missing_input_cells': {c: sum(not r['raw'][c].strip() for r in audit) for c in COLUMNS},
              'accepted_missing_email': sum(not r['email'] for r in accepted),
              'rejection_reasons': dict(Counter(issue for r in audit if r['decision'] == 'rejected' for issue in r['issues'])),
              'amount_total': f'{total:.2f}', 'amount_mean': f'{total/len(accepted):.2f}' if accepted else None,
              'category_counts': dict(Counter(r['category'] for r in accepted)),
              'status_counts': dict(Counter(r['status'] for r in accepted)), 'rules': rules}
    assert report['input_rows'] == sum(counts.values())
    return pd.DataFrame(accepted, columns=COLUMNS), audit, report


def write_workbook(path, cleaned, report):
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    summary.append(['Cleaning summary (run snapshot)'])
    summary.append(['Metric', 'Value'])
    for key in ('input_rows', 'cleaned_rows', 'duplicates_removed', 'rejected_rows', 'normalized_cells', 'accepted_missing_email'):
        summary.append([key.replace('_', ' ').title(), report[key]])
    summary.append(['Amount total (USD)', float(report['amount_total'])])
    summary.append(['Amount mean (USD)', float(report['amount_mean']) if report['amount_mean'] else None])
    summary.append([])
    summary.append(['Category', 'Accepted records'])
    for key, value in report['category_counts'].items():
        summary.append([key, value])
    sheet = wb.create_sheet('Cleaned')
    sheet.append(COLUMNS)
    for record in cleaned.to_dict('records'):
        values = [record[c] for c in COLUMNS]
        values[4] = date.fromisoformat(values[4])
        values[5] = float(values[5])
        sheet.append(values)
    for ws in wb:
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='center', horizontal='right' if cell.data_type == 'n' else 'left')
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Keep imported text literal, even if it starts with '='.
        headers = [1] if ws.title == 'Cleaned' else [2, 12]
        for n in headers:
            for cell in ws[n]:
                cell.fill = PatternFill('solid', fgColor='24364B')
                cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
                cell.alignment = Alignment(horizontal='center', vertical='center')
        for row in range(1, ws.max_row+1):
            ws.row_dimensions[row].height = 23
    summary.column_dimensions['A'].width = 43
    summary.column_dimensions['B'].width = 20
    summary['B9'].number_format = summary['B10'].number_format = '"$"#,##0.00'
    for col, width in zip('ABCDEFG', [16, 24, 34, 20, 18, 18, 18]):
        sheet.column_dimensions[col].width = width
    for row in range(2, sheet.max_row+1):
        sheet.cell(row, 5).number_format = 'yyyy-mm-dd'
        sheet.cell(row, 6).number_format = '"$"#,##0.00'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    wb.save(path)


def run(source, output, rules=None, sheet=0):
    source, output = Path(source).resolve(), Path(output).resolve()
    if output == source or output in source.parents:
        raise ValueError('Output must not contain the input file')
    if output.exists():
        raise FileExistsError('Choose a new output folder; overwrite is disabled')
    frame = read_input(source, sheet)
    cleaned, audit, report = clean(frame, rules)
    output.mkdir(parents=True)
    cleaned.to_csv(output/'cleaned.csv', index=False)
    write_workbook(output/'cleaned.xlsx', cleaned, report)
    flat = [{'source_row': r['source_row'], 'decision': r['decision'], 'issues': '; '.join(r['issues']),
             **{'raw_'+c: r['raw'][c] for c in COLUMNS}, **{'clean_'+c: r['clean'][c] for c in COLUMNS}} for r in audit]
    audit_frame = pd.DataFrame(flat, columns=['source_row', 'decision', 'issues']+['raw_'+c for c in COLUMNS]+['clean_'+c for c in COLUMNS])
    audit_frame.to_csv(output/'before_after.csv', index=False)
    audit_frame[audit_frame['decision'] == 'rejected'].to_csv(output/'rejected.csv', index=False)
    (output/'report.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    lines = ['# Cleaning summary', '', *[f'{k}: {report[k]}' for k in ('input_rows','cleaned_rows','duplicates_removed','rejected_rows','normalized_cells','accepted_missing_email','amount_total','amount_mean')], '',
             'Amounts are USD; statistics use accepted records only.', '', '## Rejection reasons',
             *[f'- {k}: {v}' for k,v in report['rejection_reasons'].items()], '', '## Category counts',
             *[f'- {k}: {v}' for k,v in report['category_counts'].items()], '', '## Rejected source rows (header is row 1)',
             *[f"- Row {r['source_row']}, {r['raw']['order_id']}: {', '.join(r['issues'])}" for r in audit if r['decision'] == 'rejected'], '',
             'See before_after.csv for every row, normalized value and disposition. Optional missing emails remain blank. No missing values are invented.']
    (output/'report.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--rules', type=Path, help='JSON category_aliases, status_aliases, date_formats overrides')
    parser.add_argument('--sheet', default=0, help='XLSX sheet name; default first sheet')
    args = parser.parse_args()
    rules = {**DEFAULT_RULES, **json.loads(args.rules.read_text())} if args.rules else DEFAULT_RULES
    print(json.dumps(run(args.input, args.output, rules, args.sheet), indent=2))
